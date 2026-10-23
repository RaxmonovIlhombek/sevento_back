"""Hisobotlar: reyestr va ruxsat, JSON natija, PDF/Excel, fon vazifasi (report.ready), filtrlar."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone

from apps.reports import services
from apps.reports.models import ReportJob
from apps.reports.registry import BaseReport, register_report
from apps.reports.types import Column
from core.models import OutboxEvent

pytestmark = pytest.mark.django_db

D = Decimal
URL = "/api/v1/reports/"


def get(client, key, **params):
    response = client.get(f"{URL}{key}/", params)
    assert response.status_code == 200, response.json()
    return response.json()


def test_registry_by_role(owner_client, storekeeper, accountant, seller_client, as_user):
    owner_keys = {r["key"] for r in owner_client.get(URL).json()}
    assert {"sales_summary", "stock_balance", "debts_aging", "purchases"} <= owner_keys
    store_keys = {r["key"] for r in as_user(storekeeper).get(URL).json()}
    # omborchi: ombor va kirim hisobotlari; moliyaviy (savdo, kassa, qarz) — yo'q
    assert "stock_balance" in store_keys and "purchases" in store_keys
    assert not store_keys & {"sales_summary", "cash_flow", "debts"}
    assert seller_client.get(URL).status_code == 403
    info = next(r for r in owner_client.get(URL).json() if r["key"] == "sales_summary")
    assert info["category_label"] == "Savdo" and info["formats"] == ["pdf", "xlsx"]
    assert [f["name"] for f in info["filters"]] == ["date_from", "date_to", "branch", "group_by"]
    group = info["filters"][-1]
    assert (
        group["default"] == "day"
        and {"value": "seller", "label": "Sotuvchilar"} in group["choices"]
    )
    assert as_user(storekeeper).get(f"{URL}sales_summary/").status_code == 403


def test_sales_summary(owner_client, world, branch):
    data = get(owner_client, "sales_summary")
    assert data["title"] == "Savdo hisoboti" and "Barcha nuqtalar" in data["subtitle"]
    row = data["rows"][0]
    # savdo 80 000 so'm = 6,40 USD; qaytarish 10 000 = 0,80; tannarx 8 × 0,40 = 3,20, qaytgan 0,40
    assert (row["count"], row["revenue"], row["returns"], row["net"]) == (2, "6.40", "0.80", "5.60")
    assert (row["cost"], row["profit"], row["margin"]) == ("2.80", "2.80", "50.0")
    assert data["totals"]["profit"] == "2.80"
    by_branch = get(owner_client, "sales_summary", group_by="branch", branch=branch.id)
    assert by_branch["rows"][0]["group"] == "Toshkent 1" and "Toshkent 1" in by_branch["subtitle"]


def test_sales_by_product(owner_client, world):
    row = get(owner_client, "sales_by_product")["rows"][0]
    assert (row["name"], row["sold"], row["returned"], row["net_qty"]) == (
        "Cola 1L",
        "8.000",
        "1.000",
        "7.000",
    )
    assert (row["revenue"], row["cost"], row["profit"]) == ("5.60", "2.80", "2.80")
    by_category = get(owner_client, "sales_by_product", group_by="category")["rows"][0]
    assert by_category["name"] == "Kategoriyasiz"


def test_inventory_reports(owner_client, world):
    stock = get(owner_client, "stock_balance")
    # 20 − 8 sotildi + 1 qaytdi − 2 hisobdan chiqdi = 11
    assert stock["rows"][0]["quantity"] == "11.000" and stock["rows"][0]["value"] == "4.40"
    turnover = get(owner_client, "stock_turnover")["rows"][0]
    assert (turnover["opening"], turnover["purchases"], turnover["sales"]) == (
        "0.000",
        "20.000",
        "-7.000",
    )
    assert (turnover["adjustments"], turnover["closing"]) == ("-2.000", "11.000")
    write_offs = get(owner_client, "write_offs")
    assert write_offs["rows"][0]["reason"] == "Buzilgan / shikastlangan"
    assert write_offs["totals"]["cost"] == "0.80"
    movements = get(owner_client, "stock_movements")
    assert movements["heavy"] is True and movements["row_count"] == 5


def test_cost_columns_need_capability(roles, storekeeper, as_user, world):
    roles["storekeeper"].capabilities = [
        c for c in roles["storekeeper"].capabilities if c != "inventory.view_cost"
    ]
    roles["storekeeper"].save()
    from apps.accounts.models import User

    client = as_user(User.objects.get(pk=storekeeper.pk))
    data = get(client, "stock_balance")
    assert "value" not in {c["key"] for c in data["columns"]}
    assert "value" not in data["rows"][0] and data["totals"] is None


def test_finance_reports(owner_client, world):
    debts = get(owner_client, "debts")
    assert debts["rows"][0]["customer"] == "Aliyev Vali"
    assert (
        debts["rows"][0]["remaining"] == "30000.0000"
        and debts["totals"]["remaining_base"] == "2.40"
    )
    aging = get(owner_client, "debts_aging")["rows"][0]
    assert aging["current"] == "2.40" and aging["total"] == "2.40"
    cash = get(owner_client, "cash_flow")["rows"][0]
    # 50 000 kirim, 10 000 qaytarildi
    assert (cash["opening"], cash["incoming"], cash["outgoing"], cash["closing"]) == (
        "0",
        "50000.0000",
        "10000.0000",
        "40000.0000",
    )
    payments = get(owner_client, "payments_summary")["rows"][0]
    assert (payments["net"], payments["net_base"]) == ("40000.0000", "3.20")
    returns = get(owner_client, "returns")["rows"][0]
    assert returns["refund"] == "10000.0000" and returns["reason"] == "Sifatsiz"
    purchases = get(owner_client, "purchases")["rows"][0]
    assert purchases["total"] == "8.00"


def test_params_validation(owner_client, seller, as_user, branch2):
    bad = owner_client.get(f"{URL}sales_summary/", {"date_from": "2026-13-01"})
    assert bad.status_code == 400 and bad.json()["error"]["details"][0]["field"] == "date_from"
    rng = owner_client.get(
        f"{URL}sales_summary/", {"date_from": "2026-10-10", "date_to": "2026-10-01"}
    )
    assert rng.json()["error"]["details"][0]["code"] == "invalid_range"
    choice = owner_client.get(f"{URL}sales_summary/", {"group_by": "planet"})
    assert choice.json()["error"]["details"][0]["code"] == "invalid_choice"
    assert owner_client.get(f"{URL}nope/").status_code == 404


def test_alien_branch_forbidden(storekeeper, as_user, branch2):
    response = as_user(storekeeper).get(f"{URL}stock_balance/", {"branch": branch2.id})
    assert response.status_code == 403 and response.json()["error"]["code"] == "branch_forbidden"


def test_pdf_export(owner_client, world):
    response = owner_client.post(
        f"{URL}sales_summary/export/", {"format": "pdf", "params": {}}, format="json"
    )
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert 'filename="sales_summary-' in response["Content-Disposition"]
    event = OutboxEvent.objects.get(name="report.exported")
    assert event.payload["report_key"] == "sales_summary" and event.payload["rows"] == 1


def test_xlsx_export(owner_client, world):
    from openpyxl import load_workbook

    response = owner_client.post(
        f"{URL}sales_by_product/export/", {"format": "xlsx", "params": {}}, format="json"
    )
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content)).active
    assert ws["A1"].value == "Tovarlar bo'yicha savdo va foyda"
    header = [c.value for c in ws[5]]
    assert header[:4] == ["Tovar", "Artikul", "Birlik", "Sotildi"]
    assert ws["A6"].value == "Cola 1L" and ws["D6"].value == 8
    assert ws["A7"].value == "Jami"
    assert ws.freeze_panes == "A6"


def test_heavy_export_runs_in_background(
    owner, owner_client, accountant, as_user, world, django_capture_on_commit_callbacks
):
    with django_capture_on_commit_callbacks(execute=True):
        response = owner_client.post(
            f"{URL}stock_movements/export/", {"format": "xlsx", "params": {}}, format="json"
        )
    assert response.status_code == 202, response.json()
    job = ReportJob.objects.get(pk=response.json()["id"])
    assert job.status == "ready" and job.rows_count == 5 and job.file
    listed = owner_client.get(f"{URL}jobs/").json()
    assert listed[0]["download_url"] == f"/api/v1/reports/jobs/{job.pk}/download/"
    download = owner_client.get(listed[0]["download_url"])
    assert download.status_code == 200 and b"".join(download.streaming_content)[:2] == b"PK"
    ready = OutboxEvent.objects.get(name="report.ready")
    assert ready.payload["user_id"] == owner.pk and ready.payload["format"] == "xlsx"
    # boshqa foydalanuvchi faylni ko'rmaydi
    assert as_user(accountant).get(listed[0]["download_url"]).status_code == 404

    from apps.realtime.router import resolve
    from core.events import Event

    groups = {o.group for o in resolve(Event.from_outbox(ready))}
    assert groups == {f"user_{owner.pk}"}


def test_long_period_goes_background(owner_client, world, django_capture_on_commit_callbacks):
    today = timezone.localdate()
    params = {"date_from": (today - timedelta(days=60)).isoformat(), "date_to": today.isoformat()}
    with django_capture_on_commit_callbacks(execute=True):
        response = owner_client.post(
            f"{URL}sales_summary/export/", {"format": "pdf", "params": params}, format="json"
        )
    assert response.status_code == 202
    assert ReportJob.objects.get().status == "ready"


def test_failed_job_reports_error(owner, world, monkeypatch, django_capture_on_commit_callbacks):
    job = ReportJob.objects.create(
        report_key="sales_summary", title="Savdo hisoboti", format="pdf", requested_by=owner
    )

    def boom(*args, **kwargs):
        raise RuntimeError("shrift topilmadi")

    monkeypatch.setattr(services, "_render", boom)
    with django_capture_on_commit_callbacks(execute=True):
        services.run_job(job.pk)
    job.refresh_from_db()
    assert job.status == "failed" and job.error == "shrift topilmadi"
    assert OutboxEvent.objects.get(name="report.failed").payload["error"] == "shrift topilmadi"


def test_cleanup_old_jobs(owner, world, owner_client, django_capture_on_commit_callbacks):
    with django_capture_on_commit_callbacks(execute=True):
        owner_client.post(f"{URL}stock_movements/export/", {"format": "pdf"}, format="json")
    job = ReportJob.objects.get()
    path = job.file.path
    assert services.cleanup_old_jobs(now=timezone.now() + timedelta(days=8)) == 1
    import os

    assert not os.path.exists(path) and not ReportJob.objects.exists()


def test_new_report_plugs_in_without_changes(owner_client):
    """Reyestr: yangi klass — mavjud view/render o'zgarmaydi."""

    @register_report
    class PingReport(BaseReport):
        key = "ping_test"
        title = "Sinov"
        category = "finance"

        def columns(self, params):
            return [Column("name", "Nomi"), Column("value", "Qiymat", "money", total=True)]

        def rows(self, params):
            return [{"name": "a", "value": D("1.5")}, {"name": "b", "value": D("2")}]

    try:
        data = get(owner_client, "ping_test")
        assert data["totals"] == {"value": "3.5"}
        pdf = owner_client.post(f"{URL}ping_test/export/", {"format": "pdf"}, format="json")
        assert pdf.status_code == 200
    finally:
        from apps.reports import registry

        registry._REPORTS.pop("ping_test", None)
