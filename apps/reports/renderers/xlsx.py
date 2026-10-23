"""Excel (openpyxl): sarlavha, filtr tavsifi, ustun formatlari, jami qatori, muzlatilgan qator."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from .base import RenderedReport, decimals_for

HEADER_FILL = "E2E8F0"


def _number_format(column, row) -> str:
    places = decimals_for(column, row)
    base = "#,##0" + ("." + "0" * places if places > 0 else "")
    return base + '"%"' if column.type == "percent" else base


class XlsxRenderer:
    format = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def render(self, report: RenderedReport) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        result = report.result
        columns = result.columns
        wb = Workbook()
        ws = wb.active
        ws.title = report.title[:31]
        ws.append([report.title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([report.subtitle])
        generated = timezone.localtime(report.generated_at).strftime("%d.%m.%Y %H:%M")
        ws.append([f"Tayyorladi: {report.generated_by} · {generated}"])
        ws.append([])
        ws.append([c.label for c in columns])
        header_row = ws.max_row
        bold = Font(bold=True)
        for index, column in enumerate(columns, start=1):
            cell = ws.cell(row=header_row, column=index)
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(
                horizontal="right" if column.numeric else "left", vertical="center", wrap_text=True
            )
        widths = [max(len(c.label), c.width or 8) for c in columns]
        for row in result.rows:
            values = []
            for index, column in enumerate(columns):
                value = row.get(column.key)
                if isinstance(value, Decimal):
                    value = float(value) if column.numeric else str(value)
                elif isinstance(value, datetime):
                    value = timezone.localtime(value).replace(tzinfo=None)
                values.append(value)
                widths[index] = max(widths[index], min(len(str(row.get(column.key) or "")), 60))
            ws.append(values)
            current = ws.max_row
            for index, column in enumerate(columns, start=1):
                cell = ws.cell(row=current, column=index)
                if column.numeric:
                    cell.number_format = _number_format(column, row)
                elif isinstance(cell.value, datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"
        if result.totals:
            ws.append(
                [float(result.totals[c.key]) if c.key in result.totals else None for c in columns]
            )
            current = ws.max_row
            if columns and columns[0].key not in result.totals:
                ws.cell(row=current, column=1, value="Jami")
            for index, column in enumerate(columns, start=1):
                cell = ws.cell(row=current, column=index)
                cell.font = bold
                if column.numeric:
                    cell.number_format = _number_format(column, None)
        for note in result.notes:
            ws.append([note])
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width + 2
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        last_column = get_column_letter(max(len(columns), 1))
        ws.auto_filter.ref = f"A{header_row}:{last_column}{header_row + len(result.rows)}"
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
